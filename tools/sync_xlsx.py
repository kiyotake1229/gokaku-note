#!/usr/bin/env python3
"""アプリの記録（学習ログ・問題の結果・受験の予定と結果）を、学習記録.xlsx に書き込む。

使い方:
  python3 tools/sync_xlsx.py <バックアップのファイル.json>     … アプリで保存したバックアップから
  python3 tools/sync_xlsx.py --latest <フォルダ>               … フォルダの中のいちばん新しいバックアップから
  python3 tools/sync_xlsx.py --setup "<いつでも最新のリンク>"    … 自動共有の「いつでも最新」のリンクを覚えさせる（最初に1回）
  python3 tools/sync_xlsx.py --gas                             … 自動共有（Google スプレッドシート）から
  （共通）--xlsx <学習記録.xlsx の場所>  既定は study-app の1つ上のフォルダ
          --dry  書き込まずに、足す行・直す行の数だけ表示

きまり:
- アプリの記録には1件ずつIDがあり、見えない列（学習ログはL列、問題の結果はJ列）に入れる。
  同じIDの行は、もう一度足さずに中身だけ新しくする（何度動かしても行が重ならない）。
- Excel に手で書いた行は、そのまま残す。コメント欄（学習ログのG列）にすでに書いてあれば上書きしない。
- 受験の予定と結果は、アプリに書いてある項目だけを書きかえる（アプリが空の項目は Excel のまま）。
- Excel で開いているとき（~$学習記録.xlsx があるとき）は書き込まない。
- 書き込む前に、tools/work/xlsx-backups/ に控えを残す（新しい20個まで）。
- --gas は、相談相手に送る「いつでも最新」のリンク（読むための鍵だけが入っている）を使う。
  リンクは ~/.gokaku-note/xlsx_sync.json に覚えておく（共有フォルダや公開リポジトリには置かない）。
"""
import json, os, sys, glob, shutil, datetime, base64, zlib, urllib.request, urllib.parse
import openpyxl
from openpyxl.styles import Font

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
args = sys.argv[1:]
DRY = '--dry' in args
XLSX = args[args.index('--xlsx') + 1] if '--xlsx' in args else os.path.join(os.path.dirname(ROOT), '学習記録.xlsx')

# 資格の表記（学習記録.xlsx のプルダウンと、受験の表の1列目に合わせる）
SHORT = {'line-basic': 'LINE Basic', 'line-adv': 'LINE Advanced', 'gads': 'Google広告認定', 'ga4': 'GA4認定',
         'itpass': 'ITパスポート', 'genai': '生成AIパスポート', 'jstqb': 'JSTQB FL', 'sg': '情報セキュリティマネジメント'}
LONG = {'line-basic': 'LINE公式アカウント Basic', 'line-adv': 'LINE公式アカウント Advanced', 'gads': 'Google広告認定資格', 'ga4': 'GA4認定資格',
        'itpass': 'ITパスポート', 'genai': '生成AIパスポート', 'jstqb': 'JSTQB Foundation Level', 'sg': '情報セキュリティマネジメント'}
MODE = {'smart': 'おまかせ学習', 'due': '今日の復習', 'mock': '模擬試験', 'wrong': '間違えた問題', 'unsure': '迷った問題', 'new': 'まだ解いていない問題',
        'mark': '印をつけた問題', 'cat': '分野別', 'past': '公開問題', 'pastexam': '公開問題（本番形式）', 'drill': '計算ドリル', 'tquiz': '用語クイズ',
        'lesson': '教科書の確認問題', 'retry': '解き直し', 'search': '検索した問題', 'newscope': '新しい範囲', 'flagged': '報告した問題'}


def die(msg):
    print(msg)
    sys.exit(1)


def unwrap(j):
    """バックアップのファイル（{data: …}）、自動共有の返事（{backup: …}）、記録そのもの、のどれでも受け取る"""
    for _ in range(3):
        if isinstance(j, dict) and isinstance(j.get('logs'), list):
            return j
        if isinstance(j, dict):
            j = j.get('data') or j.get('backup') or j.get('state')
        if isinstance(j, str):
            j = json.loads(j)
    die('アプリの記録が見つかりません（バックアップのファイルか確かめてください）')


CONF = os.environ.get('GOKAKU_XLSX_CONF') or os.path.expanduser('~/.gokaku-note/xlsx_sync.json')


def parse_live(link):
    """「いつでも最新」のリンク（report.html#live=…）から、Apps Script のURLと読むための鍵を取り出す"""
    m = link.split('live=', 1)
    if len(m) < 2:
        die('「いつでも最新」のリンクではありません（report.html#live=… の形）')
    enc = m[1].split('&')[0].strip()
    raw = base64.urlsafe_b64decode(enc[1:] + '=' * (-len(enc[1:]) % 4))
    obj = json.loads((zlib.decompress(raw, -15) if enc[0] == 'z' else raw).decode('utf-8'))
    if not obj.get('u') or not obj.get('r'):
        die('リンクに、読む先の情報が入っていません')
    return obj['u'], obj['r']


def load_state():
    if '--setup' in args:
        link = args[args.index('--setup') + 1]
        parse_live(link)
        os.makedirs(os.path.dirname(CONF), exist_ok=True)
        json.dump({'link': link}, open(CONF, 'w', encoding='utf-8'))
        os.chmod(CONF, 0o600)
        print(f'リンクを覚えました：{CONF}')
    if '--gas' in args or '--setup' in args:
        if not os.path.exists(CONF):
            print('自動共有のリンクがまだ設定されていません（--setup "<いつでも最新のリンク>" で設定）。何もしません。')
            sys.exit(0)
        u, r = parse_live(json.load(open(CONF, encoding='utf-8'))['link'])
        url = u + ('&' if '?' in u else '?') + urllib.parse.urlencode({'what': 'records', 'k': r})
        with urllib.request.urlopen(url, timeout=60) as resp:
            res = json.loads(resp.read().decode('utf-8'))
        if not res.get('ok'):
            hint = '（Apps Script を新しいコードに更新してください）' if res.get('error') == 'key' else ''
            die(f'自動共有から読めませんでした：{res.get("error")}{hint}')
        at = datetime.datetime.fromtimestamp((res.get('at') or 0) / 1000).strftime('%Y-%m-%d %H:%M') if res.get('at') else '—'
        return unwrap(res.get('records')), f'自動共有（スマホから最後に届いた時刻 {at}）'
    if '--latest' in args:
        folder = args[args.index('--latest') + 1]
        files = sorted(glob.glob(os.path.join(folder, '*.json')), key=os.path.getmtime)
        if not files:
            die(f'{folder} にバックアップ（.json）がありません')
        return unwrap(json.load(open(files[-1], encoding='utf-8'))), os.path.basename(files[-1])
    paths = [a for a in args if a.endswith('.json')]
    if not paths:
        die(__doc__)
    return unwrap(json.load(open(paths[0], encoding='utf-8'))), os.path.basename(paths[0])


def last_row(ws, cols):
    """見出しの下で、指定した列のどれかに値がある最後の行"""
    n = 1
    for r in range(2, ws.max_row + 1):
        if any(ws.cell(r, c).value not in (None, '') for c in cols):
            n = r
    return n


def id_rows(ws, col):
    return {str(ws.cell(r, col).value): r for r in range(2, ws.max_row + 1) if ws.cell(r, col).value}


def prepare_id_col(ws, col, letter):
    ws.cell(1, col).value = 'アプリID（自動）'
    ws.cell(1, col).font = Font(color='999999', size=9)
    ws.column_dimensions[letter].hidden = True


def as_date(v):
    """'2026-10-06' を Excel の日付にする（並べかえ・しぼり込みができるように）"""
    try:
        return datetime.datetime.strptime(str(v)[:10], '%Y-%m-%d')
    except (TypeError, ValueError):
        return v


def put(ws, r, c, v, stat, keep_if_filled=False, date=False):
    cur = ws.cell(r, c).value
    if date and v not in (None, ''):
        v = as_date(v)
        ws.cell(r, c).number_format = 'yyyy/mm/dd'
    if v in (None, ''):
        return
    if keep_if_filled and cur not in (None, ''):
        return
    if cur != v:
        ws.cell(r, c).value = v
        stat['changed'] += 1


def main():
    print(f'== {datetime.datetime.now():%Y-%m-%d %H:%M}')
    state, src = load_state()
    if not os.path.exists(XLSX):
        die(f'{XLSX} がありません')
    lock = os.path.join(os.path.dirname(XLSX), '~$' + os.path.basename(XLSX))
    if os.path.exists(lock) and not DRY:
        die('学習記録.xlsx を Excel で開いているようです。閉じてから、もう一度動かしてください。')
    wb = openpyxl.load_workbook(XLSX)
    stat = {'logs': 0, 'results': 0, 'exams': 0, 'changed': 0}

    # ---- 学習ログ（A:日付 B:資格 C:やったこと D:時間 E:分かったこと F:聞きたいこと G:コメント ／ L:アプリID） ----
    ws = wb['学習ログ']
    prepare_id_col(ws, 12, 'L')
    have = id_rows(ws, 12)
    nxt = last_row(ws, range(1, 8)) + 1
    for x in sorted(state.get('logs') or [], key=lambda x: (x.get('date') or '', x.get('id') or '')):
        lid = str(x.get('id') or '')
        if not lid:
            continue
        vals = [x.get('date'), SHORT.get(x.get('cert'), ''), x.get('what') or '', int(x['minutes']) if str(x.get('minutes') or '').isdigit() else x.get('minutes'),
                x.get('learned') or '', x.get('ask') or '']
        r = have.get(lid)
        if r is None:
            r = nxt; nxt += 1; stat['logs'] += 1
            ws.cell(r, 12).value = lid
        for c, v in enumerate(vals, start=1):
            put(ws, r, c, v, stat, date=c == 1)
        put(ws, r, 7, x.get('comment') or '', stat, keep_if_filled=True)

    # ---- 問題の結果（A:日付 B:資格 C:何の問題か D:問題数 E:正解数 F:正答率（式） G:間違えた分野 ／ J:アプリID） ----
    ws = wb['問題の結果']
    prepare_id_col(ws, 10, 'J')
    have = id_rows(ws, 10)
    nxt = last_row(ws, [1, 2, 3, 4, 5, 7]) + 1
    cats = {}
    try:
        cats = {c['id']: {x['id']: x['name'] for x in c['categories']} for c in json.load(open(os.path.join(ROOT, 'data', 'certs.json'), encoding='utf-8'))['certs']}
    except Exception:
        pass
    for x in sorted(state.get('sessions') or [], key=lambda x: (x.get('date') or '', x.get('ts') or 0)):
        sid = str(x.get('id') or '')
        if not sid or not x.get('total'):
            continue
        weak = x.get('weak') or '、'.join(cats.get(x.get('cert'), {}).get(k, k) for k, v in (x.get('cats') or {}).items() if v and v[0] and v[1] / v[0] < 0.8)
        what = x.get('title') if x.get('mode') == 'manual' else f"アプリ：{x.get('title') or MODE.get(x.get('mode'), x.get('mode') or '')}"
        vals = [x.get('date'), SHORT.get(x.get('cert'), ''), what, x.get('total'), x.get('correct')]
        r = have.get(sid)
        if r is None:
            r = nxt; nxt += 1; stat['results'] += 1
            ws.cell(r, 10).value = sid
        for c, v in enumerate(vals, start=1):
            put(ws, r, c, v, stat, date=c == 1)
        if not ws.cell(r, 6).value:
            ws.cell(r, 6).value = f'=IF(OR(D{r}="",D{r}=0),"",E{r}/D{r})'
            ws.cell(r, 6).number_format = '0%'
        put(ws, r, 7, weak, stat)

    # ---- 受験の予定と結果（1列目の資格名の行に書く。アプリが空の項目は Excel のまま） ----
    ws = wb['受験の予定と結果']
    rows = {str(ws.cell(r, 1).value).strip(): r for r in range(2, ws.max_row + 1) if ws.cell(r, 1).value}
    for cid, e in (state.get('exams') or {}).items():
        if not isinstance(e, dict) or cid not in LONG:
            continue
        r = rows.get(LONG[cid])
        if r is None:
            r = last_row(ws, [1]) + 1
            ws.cell(r, 1).value = LONG[cid]
            rows[LONG[cid]] = r
        before = stat['changed']
        for c, k in enumerate(['applied', 'date', 'result', 'score', 'expiry', 'proof', 'memo'], start=2):
            put(ws, r, c, e.get(k), stat, date=k in ('applied', 'date', 'expiry') and len(str(e.get(k) or '')) == 10)
        stat['exams'] += stat['changed'] > before

    print(f'読み込んだ記録：{src}')
    print(f'学習ログ 新しく{stat["logs"]}行 ／ 問題の結果 新しく{stat["results"]}行 ／ 受験の予定と結果 {stat["exams"]}資格を更新 ／ 書きかえたセル {stat["changed"]}')
    if DRY:
        print('（--dry なので書き込んでいません）')
        return
    if not (stat['logs'] or stat['results'] or stat['changed']):
        print('変わったところはありません。')
        return
    bdir = os.path.join(ROOT, 'tools', 'work', 'xlsx-backups')
    os.makedirs(bdir, exist_ok=True)
    shutil.copy2(XLSX, os.path.join(bdir, f'学習記録_{datetime.datetime.now():%Y%m%d-%H%M%S}.xlsx'))
    for old in sorted(glob.glob(os.path.join(bdir, '学習記録_*.xlsx')))[:-20]:
        os.remove(old)
    tmp = XLSX + '.tmp.xlsx'
    wb.save(tmp)
    os.replace(tmp, XLSX)
    print(f'書き込みました：{XLSX}')


if __name__ == '__main__':
    main()
